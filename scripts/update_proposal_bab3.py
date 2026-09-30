import os
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_styled_paragraph(doc, text="", style='Normal', space_before=0, space_after=6, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    p.alignment = align
    if text:
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_heading_1(doc, text):
    p = create_styled_paragraph(doc, space_before=12, space_after=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_heading_2(doc, text):
    p = create_styled_paragraph(doc, space_before=12, space_after=6, align=WD_ALIGN_PARAGRAPH.LEFT)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_heading_3(doc, text):
    p = create_styled_paragraph(doc, space_before=8, space_after=4, align=WD_ALIGN_PARAGRAPH.LEFT)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_bullet_item(doc, label, text):
    p = create_styled_paragraph(doc, space_before=2, space_after=4, align=WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.left_indent = Inches(0.25)
    r1 = p.add_run(f"•  {label}: ")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(12)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0, 0, 0)
    
    r2 = p.add_run(text)
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(12)
    r2.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_body_paragraph(doc, text):
    p = create_styled_paragraph(doc, text=text, space_before=0, space_after=6, align=WD_ALIGN_PARAGRAPH.LEFT)
    return p

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>
                <w:left w:val="none"/>
                <w:right w:val="none"/>
                <w:insideH w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
                <w:insideV w:val="none"/>
            </w:tblBorders>
        ''')
        tblPr[0].append(borders)

def build_use_case_table(doc):
    headers = ["ID Use Case", "Nama Use Case", "Aktor Utama", "Deskripsi Ringkas"]
    data = [
        ("UC-01", "Autentikasi Pengguna", "Semua Aktor", "Pengguna masuk ke dalam sistem dengan kredensial aman sesuai hak akses peran (Admin dan Verifikator)."),
        ("UC-02", "Unggah & Pemrosesan Dokumen Naskah Dinas", "Administrator & Verifikator", "Mengunggah berkas PDF/Scan bukti inovasi, memvalidasi format dan batas ukuran, menghitung hash SHA-256 berkas, dan memicu pipeline IDP."),
        ("UC-03", "Ekstraksi Otomatis melalui Pipeline IDP 5 Tahap", "Sistem (IDP Engine)", "Memproses dokumen melalui 5 tahapan: Fast Inspector, CV Preprocessing 300 DPI, Gemini Vision Extraction, Deteksi QR TTE PyZBar, dan Post-Validation Self-Healing."),
        ("UC-04", "Meninjau dan Menetapkan Keputusan Hasil Ekstraksi (HITL)", "Verifikator BRIDA", "Meninjau hasil ekstraksi metadata, memeriksa dokumen yang berstatus NEEDS_REVIEW (< 70%), dan menetapkan status keputusan parameter (Lolos/Revisi/Tolak)."),
        ("UC-05", "Mengonfirmasi Indikator dan Validasi Dokumen", "Verifikator BRIDA", "Mengunci keputusan seluruh parameter dalam satu indikator sebagai data sah dan mengubah status indikator menjadi Selesai."),
        ("UC-06", "Mengonfirmasi Final Evaluasi", "Verifikator BRIDA", "Mengesahkan hasil evaluasi keseluruhan usulan inovasi setelah seluruh indikator tuntas dikonfirmasi, mengunci jejak audit dan stempel waktu pengesahan."),
        ("UC-07", "Pembuatan Rekapitulasi & Ekspor Laporan", "Administrator", "Menghasilkan laporan rekapitulasi evaluasi indikator dan profil naskah dinas dalam format PDF/Excel."),
        ("UC-08", "Pengujian Benchmark Akurasi Multi-Iterasi", "Administrator & Developer", "Menjalankan pengujian performa ekstraksi lintas iterasi (Baseline vs CV Heuristics vs Full Multimodal AI IDP) dan menyimpan riwayat metrik evaluasi ke basis data.")
    ]
    
    table = doc.add_table(rows=len(data) + 1, cols=4)
    table.alignment = docx.enum.table.WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    
    col_widths = [Inches(1.0), Inches(2.2), Inches(1.5), Inches(2.5)]
    
    # Header row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "F0F0F0")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=150, right=150)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 0, 0)
            
    # Data rows
    for r_idx, row_data in enumerate(data):
        row_cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = val
            set_cell_margins(row_cells[c_idx], top=100, bottom=100, left=150, right=150)
            p = row_cells[c_idx].paragraphs[0]
            if c_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(10.5)
                run.font.color.rgb = RGBColor(0, 0, 0)
                if c_idx == 0:
                    run.font.bold = True
                    
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width
            
    return table

def add_use_case_detail(doc, uc_id, uc_name, actor, precondition, trigger, basic_flow_steps, alt_flow_steps, postcondition):
    add_heading_3(doc, f"Use Case {uc_id}: {uc_name}")
    
    p = create_styled_paragraph(doc, space_before=2, space_after=3)
    r1 = p.add_run("Aktor: ")
    r1.font.name = 'Times New Roman'; r1.font.bold = True; r1.font.size = Pt(12)
    r2 = p.add_run(actor)
    r2.font.name = 'Times New Roman'; r2.font.size = Pt(12)
    
    p = create_styled_paragraph(doc, space_before=2, space_after=3)
    r1 = p.add_run("Kondisi Awal (Precondition): ")
    r1.font.name = 'Times New Roman'; r1.font.bold = True; r1.font.size = Pt(12)
    r2 = p.add_run(precondition)
    r2.font.name = 'Times New Roman'; r2.font.size = Pt(12)
    
    p = create_styled_paragraph(doc, space_before=2, space_after=3)
    r1 = p.add_run("Pemicu (Trigger): ")
    r1.font.name = 'Times New Roman'; r1.font.bold = True; r1.font.size = Pt(12)
    r2 = p.add_run(trigger)
    r2.font.name = 'Times New Roman'; r2.font.size = Pt(12)
    
    p = create_styled_paragraph(doc, space_before=3, space_after=2)
    r1 = p.add_run("Alur Utama (Basic Flow):")
    r1.font.name = 'Times New Roman'; r1.font.bold = True; r1.font.size = Pt(12)
    
    for idx, step in enumerate(basic_flow_steps):
        p = create_styled_paragraph(doc, space_before=1, space_after=2)
        p.paragraph_format.left_indent = Inches(0.25)
        r = p.add_run(f"{idx+1}. {step}")
        r.font.name = 'Times New Roman'; r.font.size = Pt(12)
        
    p = create_styled_paragraph(doc, space_before=3, space_after=2)
    r1 = p.add_run("Alur Alternatif (Alternative Flow):")
    r1.font.name = 'Times New Roman'; r1.font.bold = True; r1.font.size = Pt(12)
    
    for step in alt_flow_steps:
        p = create_styled_paragraph(doc, space_before=1, space_after=2)
        p.paragraph_format.left_indent = Inches(0.25)
        r = p.add_run(step)
        r.font.name = 'Times New Roman'; r.font.size = Pt(12)
        
    p = create_styled_paragraph(doc, space_before=2, space_after=8)
    r1 = p.add_run("Kondisi Akhir (Postcondition): ")
    r1.font.name = 'Times New Roman'; r1.font.bold = True; r1.font.size = Pt(12)
    r2 = p.add_run(postcondition)
    r2.font.name = 'Times New Roman'; r2.font.size = Pt(12)

def main():
    backup_file = '1Proposal RPL_backup.docx'
    target_file = '1Proposal RPL.docx'
    
    if not os.path.exists(backup_file):
        raise FileNotFoundError(f"Backup file {backup_file} not found!")
        
    doc = docx.Document(backup_file)
    body = doc._body._element
    
    # Identify indices
    bab3_idx = None
    lampiran_idx = None
    for i, child in enumerate(body):
        texts = [n.text for n in child.iter() if n.tag.endswith('}t') and n.text]
        txt = ''.join(texts).strip()
        if 'BAB 3' in txt and bab3_idx is None:
            bab3_idx = i
        if 'LAMPIRAN' in txt and lampiran_idx is None:
            lampiran_idx = i
            
    print(f"Old BAB 3 starts at child index {bab3_idx}, LAMPIRAN starts at child index {lampiran_idx}")
    
    # Reference element for insertion
    lampiran_el = body[lampiran_idx]
    
    # Remove all old BAB 3 elements
    for idx in range(lampiran_idx - 1, bab3_idx - 1, -1):
        el_to_remove = body[idx]
        body.remove(el_to_remove)
        
    print(f"Removed {lampiran_idx - bab3_idx} old elements from BAB 3.")
    
    # Create helper dummy document to build new elements
    temp_doc = docx.Document()
    
    # --- BAB 3 HEADING ---
    add_heading_1(temp_doc, "BAB 3 – SPESIFIKASI KEBUTUHAN SPESIFIK")
    
    # --- 3.1 ANTARMUKA EKSTERNAL ---
    add_heading_2(temp_doc, "3.1 Kebutuhan Antarmuka Eksternal")
    
    add_heading_3(temp_doc, "3.1.1 Antarmuka Pengguna (User Interface)")
    add_body_paragraph(temp_doc, "Antarmuka pengguna SIP-BRIDA IDP Engine dirancang berbasis web yang modern, intuitif, dan responsif menggunakan Tailwind CSS dengan skema warna gelap (slate-900 / dark mode) guna meminimalkan kelelahan mata bagi verifikator dalam memeriksa naskah dinas dalam jumlah banyak. Tampilan sistem berpusat pada pemrosesan cerdas dokumen naskah dinas inovasi daerah dan dasbor verifikasi manusia (Human-in-the-Loop), yang terdiri atas beberapa komponen utama:")
    add_bullet_item(temp_doc, "Header & Navigasi", "Memuat identitas sistem (\"Sturdy IDP Engine - PostgreSQL Engine\"), tombol aksi cepat \"Run Benchmark\" untuk memicu pengujian performa, tombol tautan \"Database GUI\" untuk membuka pengelola basis data, dan tautan \"API Docs\" untuk mengakses dokumentasi interaktif Swagger UI (/docs).")
    add_bullet_item(temp_doc, "Kartu Ringkasan Metrik Statistik (Dashboard Cards)", "Menyajikan 4 ringkasan analitik secara real-time: (1) Total Dokumen yang tersimpan dalam basis data; (2) Tingkat Akurasi sistem secara keseluruhan berdasarkan pengujian benchmark terkini; (3) TTE / QR Terverifikasi yang mencatat jumlah dokumen dengan tanda tangan elektronik resmi (BSrE, Srikandi ANRI, Sigap BRIDA); serta (4) Dokumen yang Memerlukan Review Manual (dokumen dengan tingkat keyakinan < 70% atau field esensial kosong).")
    add_bullet_item(temp_doc, "Panel Komparasi Benchmark Multi-Iterasi", "Menampilkan tabel evaluasi akurasi sistem lintas iterasi arsitektur (Iterasi 1: Baseline Naive, Iterasi 2: CV + Heuristik + Self-Healing Regex, dan Iterasi 3: Full Multimodal AI IDP) yang memaparkan nilai presisi, recall, dan F1-score per field (Nomor Surat, Tanggal ISO, NIP 18-digit, QR/TTE) serta persentase akurasi keseluruhan.")
    add_bullet_item(temp_doc, "Modul Unggah Dokumen PDF (Upload & Process Section)", "Area interaktif drag-and-drop untuk memilih dan mengunggah berkas naskah dinas PDF (batas kapasitas hingga 20 MB) yang dilengkapi tombol eksekusi \"Jalankan Pipeline IDP\" serta konsol Live Execution Log untuk memantau tahapan proses pipeline IDP secara real-time.")
    add_bullet_item(temp_doc, "Tabel Hasil Ekstraksi Dokumen Terdaftar", "Menampilkan daftar seluruh naskah dinas yang telah diproses dalam format tabel terstruktur, mencakup nama berkas asli, nomor surat hasil ekstraksi, tanggal dokumen berformat standar ISO-8601 (YYYY-MM-DD), NIP pejabat 18 digit, status QR/TTE dengan hyperlink langsung ke portal verifikasi keaslian resmi, Confidence Score dengan pewarnaan adaptif (hijau >= 70% dan kuning < 70%), serta badge status dokumen (VALID atau NEEDS_REVIEW).")
    add_bullet_item(temp_doc, "Dasbor Verifikasi Evaluasi (Human-in-the-Loop)", "Antarmuka evaluasi parameter inovasi daerah yang menyajikan daftar indikator penilaian di sisi kiri, panel detail indikator di bagian tengah yang memuat nama parameter, hasil rekomendasi AI, skor keyakinan, dan pilihan keputusan reviewer (Lolos, Revisi, Tolak), serta tombol pengesahan bertingkat (\"Konfirmasi Indikator Ini\" dan \"Konfirmasi Final\").")
    
    add_heading_3(temp_doc, "3.1.2 Antarmuka Perangkat Lunak (Software Interface)")
    add_bullet_item(temp_doc, "Antarmuka Backend", "Dibangun menggunakan bahasa pemrograman Python 3.12 dengan framework FastAPI sebagai otak pemrosesan data yang cepat, andal, modular, dan mendukung asynchronous background processing.")
    add_bullet_item(temp_doc, "Antarmuka Basis Data", "Menggunakan basis data relasional PostgreSQL 16 Alpine yang diakses melalui SQLAlchemy ORM. Sistem juga mengimplementasikan mekanisme graceful fallback ke SQLite lokal (storage/idp_local.db) apabila koneksi ke PostgreSQL belum tersedia, memastikan pengujian mandiri dan script pemrosesan lokal tetap berjalan lancar tanpa mengalami crash.")
    add_bullet_item(temp_doc, "Antarmuka Computer Vision & Pra-pemrosesan Citra", "Mengintegrasikan library pypdfium2 untuk rasterisasi halaman PDF ke dalam citra resolusi tinggi 300 DPI, serta OpenCV (cv2) untuk merapikan posisi kertas yang miring (Auto-Deskew menggunakan kontur bounding box cv2.minAreaRect), penguatan kontras adaptif (Contrast Limited Adaptive Histogram Equalization / CLAHE), dan pembersihan bayangan hasil pemindaian scanner.")
    add_bullet_item(temp_doc, "Antarmuka AI / Multimodal Vision", "Terhubung dengan Google Gemini AI (Gemini Flash Vision) melalui Google GenAI SDK dengan konfigurasi skema keluaran JSON ketat (Strict Structured JSON Output) untuk mengenali entitas dokumen penting (nomor surat, instansi, perihal, tanggal, nama pejabat, jabatan, NIP, stempel, dan tanda tangan). Sistem dilengkapi Smart Layout & Heuristic Extractor sebagai cadangan luring (offline fallback) saat koneksi internet terputus atau limitasi kuota API tercapai.")
    add_bullet_item(temp_doc, "Antarmuka Deteksi & Validasi TTE", "Memanfaatkan library PyZBar untuk mendeteksi dan mendekode kode QR tanda tangan elektronik (TTE) resmi bersertifikat Balai Sertifikasi Elektronik (BSrE), Srikandi ANRI, maupun Sigap BRIDA, guna menjamin keaslian naskah dinas yang diverifikasi.")
    add_bullet_item(temp_doc, "Antarmuka Kontainerisasi & Deployment", "Dikemas menggunakan Docker dan Docker Compose (Dockerfile dan docker-compose.yml) yang memisahkan kontainer aplikasi (idp_engine) dan kontainer basis data (idp_postgres) sehingga siap dipublikasikan ke Docker Hub dan dipasang pada server produksi secara portabel tanpa kendala teknis.")
    
    add_heading_3(temp_doc, "3.1.3 Antarmuka Komunikasi")
    add_body_paragraph(temp_doc, "Sistem menggunakan protokol HTTPS dengan enkripsi TLS 1.3 untuk mengamankan pertukaran data antara browser pengguna dan peladen FastAPI. Seluruh pertukaran data antar-modul dan API client-server diformat menggunakan standar JSON (JavaScript Object Notation). FastAPI secara otomatis menyediakan dokumentasi antarmuka OpenAPI/Swagger pada endpoint /docs yang dapat diuji secara interaktif.")
    
    # --- 3.2 KEBUTUHAN FUNGSIONAL ---
    add_heading_2(temp_doc, "3.2 Kebutuhan Fungsional (Use Cases)")
    add_body_paragraph(temp_doc, "Berikut merupakan pemetaan kebutuhan fungsional SIP-BRIDA IDP Engine yang disusun berdasarkan alur kerja Intelligent Document Processing dan verifikasi Human-in-the-Loop:")
    
    # Add Table
    tbl = build_use_case_table(temp_doc)
    
    add_heading_2(temp_doc, "Rincian Spesifikasi Use Case Inti")
    
    # UC-01
    add_use_case_detail(
        temp_doc,
        uc_id="UC-01",
        uc_name="Autentikasi Pengguna",
        actor="Semua Aktor (Administrator, Verifikator BRIDA)",
        precondition="Pengguna membuka antarmuka web SIP-BRIDA dan data akun telah terdaftar di dalam basis data.",
        trigger="Pengguna memasukkan kredensial login (username/email dan kata sandi) lalu menekan tombol \"Masuk\".",
        basic_flow_steps=[
            "Pengguna mengisi form login dengan kredensial yang sah.",
            "Sistem memvalidasi kecocokan username dan memverifikasi kata sandi terhadap hash standar bcrypt pada basis data.",
            "Jika terverifikasi, sistem menginisiasi sesi aman atau token otorisasi dan mengarahkan pengguna ke dasbor sesuai peran hak aksesnya."
        ],
        alt_flow_steps=[
            "2a. Kredensial tidak cocok: Sistem menolak akses, menampilkan pesan kesalahan \"Kredensial tidak valid\", dan mencatat upaya login yang gagal."
        ],
        postcondition="Pengguna berhasil masuk ke dalam sistem dengan sesi dan wewenang hak akses yang sesuai."
    )
    
    # UC-02
    add_use_case_detail(
        temp_doc,
        uc_id="UC-02",
        uc_name="Unggah dan Pemrosesan Dokumen Naskah Dinas",
        actor="Administrator, Verifikator BRIDA",
        precondition="Pengguna telah login ke dalam sistem dan berkas naskah dinas dalam format PDF telah disiapkan pada perangkat pengguna.",
        trigger="Pengguna memilih berkas PDF melalui modul unggah lalu menekan tombol \"Jalankan Pipeline IDP\".",
        basic_flow_steps=[
            "Pengguna memilih file dokumen naskah dinas PDF melalui antarmuka drag-and-drop.",
            "Sistem memvalidasi tipe konten berkas (harus format .pdf) dan memastikan ukuran berkas tidak melebihi 20 MB.",
            "Sistem menghitung nilai checksum hash SHA-256 berkas untuk mendeteksi duplikasi dokumen dan menjaga integritas naskah dinas.",
            "Sistem menyimpan berkas fisik pada direktori terproteksi (storage/), mencatat metadata dokumen baru ke tabel documents dengan status \"pending\", dan secara otomatis memicu eksekusi pipeline IDP (UC-03)."
        ],
        alt_flow_steps=[
            "2a. Format berkas bukan PDF: Sistem menolak unggahan dan menampilkan notifikasi kesalahan \"Hanya berkas PDF yang didukung\".",
            "2b. Ukuran berkas melebihi 20 MB: Sistem menolak unggahan dan menampilkan peringatan kapasitas berkas melampaui batas maksimal."
        ],
        postcondition="Berkas PDF berhasil tersimpan dan tercatat di database dengan status \"pending\", serta pemrosesan pipeline IDP dimulai secara otomatis."
    )
    
    # UC-03
    add_use_case_detail(
        temp_doc,
        uc_id="UC-03",
        uc_name="Ekstraksi Otomatis melalui Pipeline IDP 5 Tahap",
        actor="Sistem (IDP Engine)",
        precondition="Berkas naskah dinas telah tersimpan di tabel documents dengan status \"pending\" (hasil UC-02).",
        trigger="Pemanggilan asinkron fungsi pemrosesan dokumen oleh PipelineRunner.",
        basic_flow_steps=[
            "Tahap 1 (Fast Inspector): Sistem menganalisis kepadatan teks per halaman untuk mengklasifikasikan dokumen sebagai Digital Native atau Scanned Document.",
            "Tahap 2 (CV Preprocessing): Sistem merasterisasi halaman PDF menjadi citra resolusi tinggi 300 DPI via pypdfium2, mendeteksi kemiringan sudut kontur kertas via cv2.minAreaRect untuk Auto-Deskew, menerapkan penguatan kontras adaptif CLAHE, dan membersihkan bayangan scanner.",
            "Tahap 3 (Semantic Vision Extraction): Citra dokumen dikirim ke Google Gemini AI (Multimodal Vision API) dengan skema keluaran JSON terstruktur ketat untuk mengekstraksi nomor surat, instansi, perihal, tanggal, nama pejabat, jabatan, NIP, serta status stempel dan tanda tangan.",
            "Tahap 4 (Heuristic QR Detection): Sistem memanfaatkan library PyZBar untuk mencari dan mendekode kode QR tanda tangan elektronik (TTE BSrE, Srikandi ANRI, Sigap BRIDA) dan mengekstraksi URL verifikasi resmi.",
            "Tahap 5 (Post-Validation & Self-Healing): Sistem memvalidasi dan memperbaiki format NIP pegawai 18 digit menggunakan ekspresi reguler deterministik, menormalisasi tanggal naskah ke format baku ISO-8601 (YYYY-MM-DD), menghitung skor keyakinan gabungan (Confidence Score), dan menetapkan penanda needs_manual_review (True jika skor < 70% atau field esensial kosong).",
            "Sistem menyimpan seluruh hasil ekstraksi ke tabel document_extractions dan memperbarui status dokumen pada tabel documents menjadi \"completed\"."
        ],
        alt_flow_steps=[
            "3a. Layanan Gemini AI tidak terjangkau (offline / limit kuota): Sistem secara otomatis mengaktifkan Smart Layout & Heuristic Extractor sebagai fallback lokal, mengekstraksi data berbasis pola teks tanpa membuat sistem crash."
        ],
        postcondition="Seluruh metadata dokumen naskah dinas berhasil diekstraksi, divalidasi, dan disimpan ke tabel document_extractions, siap disajikan pada dasbor verifikasi."
    )
    
    # UC-04
    add_use_case_detail(
        temp_doc,
        uc_id="UC-04",
        uc_name="Meninjau dan Menetapkan Keputusan Hasil Ekstraksi (Human-in-the-Loop)",
        actor="Verifikator BRIDA",
        precondition="Dokumen naskah dinas telah selesai diproses oleh pipeline IDP (hasil UC-03) dan tercatat pada tabel ekstraksi.",
        trigger="Verifikator membuka daftar dokumen atau memilih indikator inovasi pada Dasbor Verifikasi Evaluasi.",
        basic_flow_steps=[
            "Sistem menampilkan rincian hasil ekstraksi dokumen (nomor surat, instansi, tanggal ISO, NIP pejabat, tautan verifikasi QR TTE, dan confidence score).",
            "Verifikator memeriksa kecocokan data, terutama pada dokumen berstatus NEEDS_REVIEW (skor < 70%).",
            "Verifikator dapat mengklik tautan QR TTE untuk membuka portal resmi (BSrE/Srikandi/Sigap) guna memastikan keaslian tanda tangan elektronik.",
            "Verifikator menetapkan keputusan terhadap parameter penilaian dengan memilih Lolos, Revisi, atau Tolak, serta menambahkan catatan evaluasi bila diperlukan.",
            "Sistem memperbarui keputusan reviewer pada basis data secara real-time (auto-save)."
        ],
        alt_flow_steps=[
            "2a. Terdapat data ekstraksi yang perlu disesuaikan: Verifikator melakukan koreksi langsung pada nilai metadata (Score Override) dan sistem mencatat riwayat perubahan tersebut dalam audit log."
        ],
        postcondition="Keputusan verifikator manusia tersimpan aman pada basis data sebagai penilaian resmi indikator."
    )
    
    # UC-05
    add_use_case_detail(
        temp_doc,
        uc_id="UC-05",
        uc_name="Mengonfirmasi Indikator dan Validasi Dokumen",
        actor="Verifikator BRIDA",
        precondition="Seluruh parameter pada indikator yang dipilih telah memiliki keputusan reviewer (hasil UC-04).",
        trigger="Verifikator menekan tombol \"Konfirmasi Indikator Ini\".",
        basic_flow_steps=[
            "Verifikator meninjau ringkasan keputusan seluruh parameter pada indikator yang sedang aktif.",
            "Verifikator (opsional) menambahkan catatan umum pada kolom catatan indikator.",
            "Verifikator menekan tombol \"Konfirmasi Indikator Ini\".",
            "Sistem memvalidasi bahwa seluruh parameter di dalam indikator telah memiliki keputusan resmi.",
            "Sistem mengunci keputusan indikator tersebut, mengubah status indikator dari \"Menunggu Konfirmasi\" menjadi \"Selesai\", serta memperbarui penghitung progres evaluasi (misal: \"X / 20 Indikator Selesai\") pada header dasbor."
        ],
        alt_flow_steps=[
            "4a. Masih ada parameter yang belum diputuskan: Sistem menolak konfirmasi, menampilkan pesan peringatan, dan menyoroti parameter yang belum diisi."
        ],
        postcondition="Status indikator terkunci sebagai \"Selesai\" dan progres penghitung indikator bertambah secara resmi."
    )
    
    # UC-06
    add_use_case_detail(
        temp_doc,
        uc_id="UC-06",
        uc_name="Mengonfirmasi Final Evaluasi",
        actor="Verifikator BRIDA",
        precondition="Seluruh indikator penilaian usulan inovasi telah berstatus \"Selesai\" (hasil UC-05 tuntas untuk seluruh indikator).",
        trigger="Verifikator menekan tombol \"Konfirmasi Final\" pada header dasbor.",
        basic_flow_steps=[
            "Sistem memvalidasi bahwa seluruh indikator telah berstatus \"Selesai\" (100% lengkap).",
            "Sistem menampilkan modal ringkasan akhir evaluasi (rekapitulasi jumlah parameter Lolos, Revisi, dan Tolak).",
            "Verifikator menekan tombol persetujuan akhir.",
            "Sistem mengunci seluruh hasil evaluasi sebagai catatan resmi, merekam identitas verifikator penanggung jawab dan stempel waktu pengesahan, serta mengubah status usulan inovasi menjadi \"Terverifikasi Final\"."
        ],
        alt_flow_steps=[
            "1a. Masih terdapat indikator yang belum selesai: Tombol \"Konfirmasi Final\" dinonaktifkan dan sistem menampilkan jumlah indikator yang masih tertunda."
        ],
        postcondition="Hasil evaluasi usulan inovasi bersifat final, berkas terkunci dari perubahan otomatis AI, dan data siap diterbitkan ke laporan rekapitulasi."
    )
    
    # UC-07
    add_use_case_detail(
        temp_doc,
        uc_id="UC-07",
        uc_name="Pembuatan Rekapitulasi & Ekspor Laporan",
        actor="Administrator",
        precondition="Terdapat minimal satu usulan inovasi berstatus \"Terverifikasi Final\" (hasil UC-06).",
        trigger="Administrator membuka menu \"Rekapitulasi Laporan\".",
        basic_flow_steps=[
            "Administrator memilih periode evaluasi atau usulan inovasi yang telah berstatus final.",
            "Sistem menampilkan ringkasan rekapitulasi nilai dan status kelengkapan dokumen pendukung.",
            "Administrator memilih format ekspor berkas yang diinginkan (PDF atau Excel) dan menekan \"Unduh Laporan\".",
            "Sistem men-generate berkas laporan dan menyediakan tautan unduhan secara instan."
        ],
        alt_flow_steps=[
            "1a. Belum ada dokumen yang berstatus \"Terverifikasi Final\": Sistem menampilkan pemberitahuan bahwa data laporan belum tersedia untuk periode terpilih."
        ],
        postcondition="Berkas laporan rekapitulasi resmi berhasil diunduh untuk kebutuhan pengarsipan dan pelaporan kepada pimpinan BRIDA."
    )
    
    # UC-08
    add_use_case_detail(
        temp_doc,
        uc_id="UC-08",
        uc_name="Pengujian Benchmark Akurasi Multi-Iterasi",
        actor="Administrator, AI Engineer / Developer",
        precondition="Dataset dokumen uji dan berkas data acuan (ground_truth.json) telah tersedia di dalam sistem.",
        trigger="Pengguna menekan tombol \"Run Benchmark\" pada navigasi dasbor atau memanggil endpoint /api/benchmark/run.",
        basic_flow_steps=[
            "Sistem menjalankan evaluasi pada Iterasi 1: Baseline (ekstraksi teks mentah tanpa bantuan Computer Vision).",
            "Sistem menjalankan evaluasi pada Iterasi 2: CV + Heuristik + Self-Healing (rasterisasi 300 DPI, auto-deskew, QR detection, dan perbaikan regex NIP/tanggal).",
            "Sistem menjalankan evaluasi pada Iterasi 3: Full Multimodal AI IDP (kombinasi 5-stage pipeline lengkap dengan Gemini Vision).",
            "Sistem menghitung metrik akurasi keseluruhan (overall accuracy) dan metrik evaluasi per field (F1-score) dengan membandingkannya terhadap ground truth.",
            "Sistem menyimpan riwayat pengujian ke tabel evaluation_runs di database PostgreSQL.",
            "Dashboard secara otomatis memperbarui tabel komparasi benchmark dan menampilkan kenaikan akurasi antar-iterasi."
        ],
        alt_flow_steps=[
            "1a. Berkas ground truth tidak ditemukan: Sistem membatalkan eksekusi benchmark dan mencatat error log terkait path berkas acuan."
        ],
        postcondition="Riwayat performa akurasi sistem tersimpan di basis data dan ditampilkan pada panel komparasi benchmark di antarmuka web."
    )
    
    # --- 3.3 KEBUTUHAN NON-FUNGSIONAL ---
    add_heading_2(temp_doc, "3.3 Kebutuhan Non-Fungsional")
    
    add_heading_3(temp_doc, "3.3.1 Performance")
    add_bullet_item(temp_doc, "Waktu Pemrosesan Pipeline", "Eksekusi pipeline IDP lengkap per dokumen PDF (rasterisasi 300 DPI, auto-deskew OpenCV, decoding QR PyZBar, ekstraksi Gemini Vision, dan post-validation) ditargetkan selesai dalam waktu rata-rata kurang dari 10 detik.")
    add_bullet_item(temp_doc, "Waktu Pemuatan Dasbor", "Dasbor Verifikasi Evaluasi berbasis Tailwind CSS harus dapat dimuat secara utuh dalam waktu maksimal 2 detik pada kondisi koneksi internet normal dengan kecepatan minimal 10 Mbps.")
    add_bullet_item(temp_doc, "Penyimpanan Real-Time (Auto-Save)", "Perubahan keputusan reviewer manusia pada setiap parameter harus tersimpan secara otomatis (real-time auto-save) dalam waktu kurang dari 1 detik tanpa memerlukan reload halaman secara keseluruhan.")
    add_bullet_item(temp_doc, "Pemrosesan Asinkron", "Seluruh proses berat pemrosesan dokumen dan evaluasi benchmark ditangani secara asinkron di backend FastAPI (menggunakan Pipeline Runner / Background Tasks) sehingga tidak memblokir antarmuka pengguna dan tidak mengganggu alur kerja pengguna lain.")
    
    add_heading_3(temp_doc, "3.3.2 Security")
    add_bullet_item(temp_doc, "Enkripsi Kredensial", "Mekanisme autentikasi wajib menggunakan enkripsi kata sandi standar industri berbasis algoritma bcrypt.")
    add_bullet_item(temp_doc, "Integritas Berkas via Hash SHA-256", "Setiap berkas PDF yang diunggah dihitung nilai hash kriptografis SHA-256 unik untuk mendeteksi duplikasi berkas secara instan serta menjamin integritas fisik dokumen dinas.")
    add_bullet_item(temp_doc, "Perlindungan Akses Berkas", "Berkas dokumen fisik yang diunggah disimpan di direktori terisolasi pada sisi server dengan proteksi pemblokiran akses langsung URL publik (direct public URL access blocked).")
    add_bullet_item(temp_doc, "Keamanan Kunci API & Konfigurasi", "Kunci akses Google Gemini API dan kredensial basis data PostgreSQL disimpan secara aman dalam berkas environment (.env) yang terisolasi dan tidak terekspos ke sisi klien.")
    add_bullet_item(temp_doc, "Role-Based Access Control (RBAC)", "Setiap akses fungsional dan endpoint API dibatasi secara ketat menggunakan otorisasi berbasis peran (Admin, Verifikator BRIDA).")
    
    add_heading_3(temp_doc, "3.3.3 Reliability & Availability")
    add_bullet_item(temp_doc, "Ketersediaan Sistem (Uptime)", "Ketersediaan sistem operasional ditargetkan minimal 99% selama masa aktif periode pelaporan inovasi daerah.")
    add_bullet_item(temp_doc, "Pencadangan Basis Data", "Tersedia mekanisme pencadangan (backup) basis data PostgreSQL secara otomatis setiap 24 jam sekali.")
    add_bullet_item(temp_doc, "Graceful AI Fallback Mode", "Apabila layanan eksternal Google Gemini API mengalami gangguan atau batas kuota terlampaui, sistem secara otomatis beralih ke Smart Layout & Heuristic Extractor luring, sehingga verifikator tetap dapat memproses berkas tanpa mengalami crash.")
    add_bullet_item(temp_doc, "Graceful Database Fallback Mode", "Apabila kontainer PostgreSQL Docker belum diaktifkan pada lingkungan lokal, sistem secara mandiri beralih ke basis data SQLite lokal (storage/idp_local.db) agar seluruh fungsi pengujian dan ekstraksi tetap berjalan mulus.")
    
    add_heading_3(temp_doc, "3.3.4 Usability")
    add_bullet_item(temp_doc, "Antarmuka Responsif & Tema Gelap", "Antarmuka dirancang modern menggunakan tema gelap (dark mode slate-900) berbasis Tailwind CSS yang nyaman dipandang saat bekerja dalam durasi panjang serta responsif pada resolusi layar minimal 1366x768 piksel.")
    add_bullet_item(temp_doc, "Visualisasi Status & Confidence Score", "Tingkat keyakinan sistem (Confidence Score) divisualisasikan dengan warna yang jelas (hijau untuk >= 70% dan kuning untuk < 70%) disertai penanda status (VALID vs NEEDS_REVIEW) untuk mempercepat pengambilan keputusan reviewer.")
    add_bullet_item(temp_doc, "Tautan Interaktif Verifikasi TTE", "Menyediakan hyperlink langsung pada hasil pembacaan QR Code yang langsung mengarahkan verifikator ke portal verifikasi keaslian naskah dinas resmi pemerintah (BSrE, Srikandi ANRI, Sigap BRIDA).")
    
    add_heading_3(temp_doc, "3.3.5 Portability & Deployment")
    add_bullet_item(temp_doc, "Standarisasi Kontainer Docker", "Seluruh arsitektur sistem dikemas secara modular menggunakan Docker dan Docker Compose (Dockerfile dan docker-compose.yml), mengisolasi service idp_engine dan database idp_postgres.")
    add_bullet_item(temp_doc, "Kesiapan Distribusi Docker Hub", "Sistem dirancang siap dibangun (build) dan di-push ke repositori Docker Hub, memungkinkan deployment cepat pada berbagai infrastruktur server tanpa dependensi sistem operasi luar selain Docker Engine.")
    add_bullet_item(temp_doc, "Isolasi Python Virtual Environment", "Untuk pengembangan lokal mandiri, seluruh dependensi sistem dipaketkan secara terisolasi di dalam Python Virtual Environment (.venv), mencegah terjadinya konflik versi pustaka.")
    
    # --- 3.4 KEBUTUHAN LOGICAL DATA STRUCTURE ---
    add_heading_2(temp_doc, "3.4 Kebutuhan Logical Data Structure")
    add_body_paragraph(temp_doc, "Struktur data logis dirancang secara relasional untuk memisahkan data berkas fisik, hasil ekstraksi otomatis kecerdasan buatan, riwayat evaluasi benchmark, serta data verifikasi akhir oleh manusia:")
    
    add_heading_3(temp_doc, "3.4.1 Rincian Kamus Data Entitas Kunci:")
    add_bullet_item(temp_doc, "Dokumen (Document / Tabel documents)", "Menyimpan identitas fisik berkas naskah dinas, mencakup id (Integer, Primary Key), original_filename (String 255), file_path (Text), file_hash (String 64, hash SHA-256 untuk deteksi duplikasi dan integritas), file_size_bytes (Integer), page_count (Integer), is_scanned (Boolean, penanda Scanned vs Digital Native), status (String 32: pending, processing, completed, failed), needs_manual_review (Boolean, penanda apakah skor < 70% atau field esensial kosong), serta created_at dan updated_at (DateTime).")
    add_bullet_item(temp_doc, "Hasil Ekstraksi Dokumen (DocumentExtraction / Tabel document_extractions)", "Menyimpan hasil ekstraksi cerdas metadata surat, mencakup id (Integer, Primary Key), document_id (Integer, Foreign Key ke documents.id), nomor_surat (String 255, Indexed), instansi (Text), perihal (Text), tanggal_surat (String 32, standar ISO-8601 YYYY-MM-DD), nama_pejabat (String 255), jabatan_pejabat (String 255), nip_pejabat (String 32, 18 digit tervalidasi), ada_stempel_basah (Boolean), ada_tanda_tangan (Boolean), verification_url (Text, URL hasil baca QR Code TTE BSrE/Srikandi/Sigap), confidence_score (Float 0.0-1.0), raw_json (JSON), dan created_at (DateTime).")
    add_bullet_item(temp_doc, "Pengujian Benchmark (EvaluationRun / Tabel evaluation_runs)", "Menyimpan riwayat metrik komparasi akurasi sistem antar-iterasi, mencakup id (Integer, Primary Key), iteration_name (String 64, nama iterasi pengujian), dataset_split (String 32), total_documents (Integer), accuracy_overall (Float, akurasi total), field_accuracies (JSON, presisi dan F1-score per field), metrics_detail (JSON, rekaman perbandingan prediksi terhadap ground truth), dan created_at (DateTime).")
    add_bullet_item(temp_doc, "Pengguna Sistem (User / Tabel users)", "Menyimpan identitas akun pengguna, kata sandi terenkripsi (bcrypt), peran sistem (Administrator, Verifikator BRIDA), dan instansi kedinasan.")
    add_bullet_item(temp_doc, "Verifikasi Final (Verifikasi_Final / Tabel verifikasi_final)", "Menyimpan keputusan resmi verifikator manusia atas berkas indikator inovasi, status penyesuaian nilai AI (score override), catatan rekomendasi, identitas verifikator penanggung jawab, dan stempel waktu pengesahan akhir (Human-in-the-Loop audit trail).")
    
    # Now, insert all elements from temp_doc before lampiran_el in original doc
    temp_body = temp_doc._body._element
    temp_children = list(temp_body)
    
    print(f"Transferring {len(temp_children)} new elements into the main document before LAMPIRAN...")
    for child in temp_children:
        lampiran_el.addprevious(child)
        
    doc.save(target_file)
    print(f"Successfully saved updated document to {target_file}!")

if __name__ == "__main__":
    main()
